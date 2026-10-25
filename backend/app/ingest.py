"""Turn uploaded files into clean text chunks."""
import io
import re

SUPPORTED = {"pdf", "docx", "txt", "md", "csv", "json", "log"}


def extract(filename: str, data: bytes) -> list[tuple[int | None, str]]:
    """Return a list of (page_number, text). Page is None for non-paged formats."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type '.{ext}'. Allowed: {', '.join(sorted(SUPPORTED))}")

    if ext == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]

    if ext == "docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return [(None, "\n".join(parts))]

    return [(None, data.decode("utf-8", errors="ignore"))]


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _split_long(paragraph: str, size: int) -> list[str]:
    """Split a paragraph longer than `size` on word boundaries."""
    words, out, cur = paragraph.split(), [], ""
    for w in words:
        if len(cur) + len(w) + 1 > size and cur:
            out.append(cur)
            cur = w
        else:
            cur = f"{cur} {w}" if cur else w
    if cur:
        out.append(cur)
    return out


def chunk_text(text: str, size: int = 1000, overlap: int = 150) -> list[str]:
    """Paragraph-aware chunking with a small character overlap between chunks."""
    text = _clean(text)
    if not text:
        return []
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    pieces: list[str] = []
    for p in paragraphs:
        pieces.extend(_split_long(p, size) if len(p) > size else [p])

    chunks, cur = [], ""
    for piece in pieces:
        if not cur:
            cur = piece
        elif len(cur) + len(piece) + 2 <= size:
            cur = f"{cur}\n\n{piece}"
        else:
            chunks.append(cur)
            tail = cur[-overlap:] if overlap > 0 else ""
            if " " in tail:
                tail = tail[tail.find(" ") + 1:]  # start the overlap on a whole word
            cur = f"{tail}\n\n{piece}" if tail else piece
    if cur:
        chunks.append(cur)
    return chunks
